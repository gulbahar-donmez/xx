"""Exam (deneme) result import + YKS-distribution-aware analysis.
Excel/CSV parsed automatically, PDF parsed when its text layer holds a subject table;
otherwise the student edits rows manually. Question counts follow the ÖSYM YKS layout."""

import csv
import io
import re
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile

from lib.auth import require_roles
from lib.db import db
from lib.points import award
from lib.levels import LEVEL_NAMES, POINTS, level_for, now_iso
from models.academy import ExamDistribution, ExamIn, ExamParse, ExamResult

router = APIRouter(prefix="/exams", tags=["exams"])
student_only = require_roles("student")
MAX_BYTES = 5 * 1024 * 1024
TR_LOWER = str.maketrans("İIÇĞÖŞÜ", "iıçğöşü")

# ÖSYM YKS soru dağılımı: (ders, soru sayısı, test grubu)
DISTRIBUTIONS: dict[str, dict] = {
    "TYT": {"label": "TYT · Temel Yeterlilik Testi", "duration": 165, "items": [
        ("Türkçe", 40, "Türkçe"), ("Tarih", 5, "Sosyal Bilimler"), ("Coğrafya", 5, "Sosyal Bilimler"), ("Felsefe", 5, "Sosyal Bilimler"),
        ("Din Kültürü", 5, "Sosyal Bilimler"), ("Matematik", 40, "Temel Matematik"), ("Fizik", 7, "Fen Bilimleri"), ("Kimya", 7, "Fen Bilimleri"), ("Biyoloji", 6, "Fen Bilimleri")]},
    "AYT-SAY": {"label": "AYT · Sayısal", "duration": 180, "items": [
        ("Matematik", 40, "Matematik"), ("Fizik", 14, "Fen Bilimleri"), ("Kimya", 13, "Fen Bilimleri"), ("Biyoloji", 13, "Fen Bilimleri")]},
    "AYT-EA": {"label": "AYT · Eşit Ağırlık", "duration": 180, "items": [
        ("Matematik", 40, "Matematik"), ("Edebiyat", 24, "Edebiyat-Sosyal 1"), ("Tarih", 10, "Edebiyat-Sosyal 1"), ("Coğrafya", 6, "Edebiyat-Sosyal 1")]},
    "AYT-SOZ": {"label": "AYT · Sözel", "duration": 180, "items": [
        ("Edebiyat", 24, "Edebiyat-Sosyal 1"), ("Tarih", 10, "Edebiyat-Sosyal 1"), ("Coğrafya", 6, "Edebiyat-Sosyal 1"), ("Tarih-2", 11, "Sosyal Bilimler 2"),
        ("Coğrafya-2", 11, "Sosyal Bilimler 2"), ("Felsefe Grubu", 12, "Sosyal Bilimler 2"), ("Din Kültürü", 6, "Sosyal Bilimler 2")]},
}
# Grup satırı yüklenirse (ör. "Fen Bilimleri 20") toplam soru sayısı
GROUP_TOTALS = {"TYT": {"Sosyal Bilimler": 20, "Fen Bilimleri": 20}}
ALIASES = {
    "türkçe": "Türkçe", "turkce": "Türkçe", "temel mat": "Matematik", "matematik": "Matematik", "mat": "Matematik", "geometri": "Matematik",
    "fizik": "Fizik", "kimya": "Kimya", "biyoloji": "Biyoloji", "tarih-2": "Tarih-2", "tarih 2": "Tarih-2", "tarih": "Tarih",
    "coğrafya-2": "Coğrafya-2", "coğrafya 2": "Coğrafya-2", "coğrafya": "Coğrafya", "cografya": "Coğrafya", "felsefe grubu": "Felsefe Grubu",
    "felsefe": "Felsefe", "din": "Din Kültürü", "sosyal": "Sosyal Bilimler", "fen": "Fen Bilimleri", "edebiyat": "Edebiyat", "türk dili": "Edebiyat",
}
LEVEL_ADVICE = {
    1: "Temel konu anlatımına dön; kolay sorularla kavramları oturt, günde 15 soru çöz.",
    2: "Konu eksiklerini kapat; orta seviye bağlam temelli sorularla günde 25 soru çöz.",
    3: "Zor sorulara geç; yanlışlarının hata kökü analizini yap ve süre tut.",
    4: "Seviyeni koru; haftada 2 deneme temposunda tarama testi yeterli.",
}


def lower(text: str) -> str:
    return str(text).translate(TR_LOWER).lower().strip()


def subject_of(text: str) -> str | None:
    value = lower(text)
    for key, name in ALIASES.items():
        if value.startswith(key):
            return name
    return None


def to_int(value) -> int:
    try:
        return max(0, int(float(str(value).replace(",", "."))))
    except (TypeError, ValueError):
        return 0


def max_for(exam_type: str, subject: str) -> int:
    for name, count, _group in DISTRIBUTIONS[exam_type]["items"]:
        if name == subject:
            return count
    return GROUP_TOTALS.get(exam_type, {}).get(subject, 0)


def analyse(row: dict, exam_type: str) -> dict:
    limit = max_for(exam_type, row["subject"])
    correct, wrong = row["correct"], row["wrong"]
    if limit and correct + wrong > limit:
        raise HTTPException(status_code=422, detail=f"{row['subject']} için doğru + yanlış {limit} soruyu geçemez ({exam_type})")
    blank = max(0, limit - correct - wrong) if limit else row["blank"]
    net = round(correct - wrong / 4, 2)
    total = limit or max(1, correct + wrong + blank)
    rate = max(0, round(net / total * 100))
    level = level_for(rate)
    tip = LEVEL_ADVICE[level]
    if wrong >= 4 and wrong > correct / 2:
        tip += " Yanlış sayın yüksek: emin olmadığın soruyu boş bırakmak netini korur."
    return {"subject": row["subject"], "correct": correct, "wrong": wrong, "blank": blank, "net": net, "max_questions": limit, "rate": rate, "level": level, "advice": f"{LEVEL_NAMES[level]} seviye · {tip}"}


def default_rows(exam_type: str) -> list[dict]:
    return [{"subject": n, "correct": 0, "wrong": 0, "blank": c, "net": 0, "max_questions": c, "rate": 0, "level": 0, "advice": ""} for n, c, _ in DISTRIBUTIONS[exam_type]["items"]]


def rows_from_table(table: list[list]) -> list[dict]:
    cols, rows = None, []
    for raw in table:
        cells = ["" if c is None else str(c) for c in raw]
        lowered = [lower(c) for c in cells]
        if cols is None and any("doğru" in c or "dogru" in c for c in lowered):
            find = lambda *keys: next((i for i, c in enumerate(lowered) if any(k in c for k in keys)), None)
            cols = {"subject": find("ders", "alan", "test") or 0, "correct": find("doğru", "dogru"), "wrong": find("yanlış", "yanlis"), "blank": find("boş", "bos")}
            continue
        if cols and cols["correct"] is not None and cols["wrong"] is not None:
            subject = subject_of(cells[cols["subject"]]) if cols["subject"] < len(cells) else None
            if subject:
                get = lambda key: to_int(cells[cols[key]]) if cols[key] is not None and cols[key] < len(cells) else 0
                rows.append({"subject": subject, "correct": get("correct"), "wrong": get("wrong"), "blank": get("blank")})
        elif len(cells) >= 3 and subject_of(cells[0]):
            nums = [to_int(c) for c in cells[1:4]] + [0, 0]
            rows.append({"subject": subject_of(cells[0]), "correct": nums[0], "wrong": nums[1], "blank": nums[2]})
    return rows


def parse_excel(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(data), data_only=True, read_only=True).active
    return rows_from_table([list(r) for r in sheet.iter_rows(values_only=True)])


def parse_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="ignore")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
    return rows_from_table(list(csv.reader(io.StringIO(text), dialect)))


def parse_pdf(data: bytes) -> list[dict]:
    from pypdf import PdfReader

    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    rows, seen = [], set()
    for match in re.finditer(r"([A-Za-zÇĞİÖŞÜçğıöşü][A-Za-zÇĞİÖŞÜçğıöşü .\-2]{1,30}?)\s*[:\-]?\s+(\d{1,3})\s+(\d{1,3})\s+(\d{1,3})", text):
        subject = subject_of(match.group(1))
        if subject and subject not in seen:
            seen.add(subject)
            rows.append({"subject": subject, "correct": int(match.group(2)), "wrong": int(match.group(3)), "blank": int(match.group(4))})
    return rows


def preview(rows: list[dict], exam_type: str) -> list[dict]:
    out = []
    for r in rows:
        try:
            out.append(analyse(r, exam_type))
        except HTTPException:
            out.append({**r, "net": round(r["correct"] - r["wrong"] / 4, 2), "max_questions": max_for(exam_type, r["subject"]), "rate": 0, "level": 0, "advice": "Soru sayısı YKS dağılımını aşıyor, kontrol et."})
    return out


@router.get("/distributions", response_model=list[ExamDistribution])
async def distributions():
    return [{"exam_type": k, "label": v["label"], "duration_minutes": v["duration"], "total_questions": sum(c for _, c, _ in v["items"]),
             "items": [{"subject": n, "questions": c, "group": g} for n, c, g in v["items"]]} for k, v in DISTRIBUTIONS.items()]


@router.post("/parse", response_model=ExamParse)
async def parse_exam(file: UploadFile = File(...), exam_type: str = Form("TYT"), user: dict = Depends(student_only)):
    if exam_type not in DISTRIBUTIONS:
        raise HTTPException(status_code=422, detail="Geçersiz sınav türü")
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="Dosya 5 MB sınırını aşıyor")
    name = (file.filename or "").lower()
    parsers = {".xlsx": ("excel", parse_excel), ".xlsm": ("excel", parse_excel), ".csv": ("csv", parse_csv), ".pdf": ("pdf", parse_pdf)}
    ext = next((e for e in parsers if name.endswith(e)), None)
    if not ext:
        raise HTTPException(status_code=415, detail="Desteklenen formatlar: .xlsx, .csv, .pdf")
    source, parser = parsers[ext]
    try:
        rows = parser(data)
    except Exception:
        rows = []
    if rows:
        return {"exam_type": exam_type, "source": source, "parsed": True, "message": f"{len(rows)} ders satırı otomatik okundu. {exam_type} soru dağılımına göre kontrol edip kaydet.", "rows": preview(rows, exam_type)}
    return {"exam_type": exam_type, "source": source, "parsed": False, "message": "Dosyadaki tablo okunamadı. Değerleri YKS dağılımına göre hazırlanan alanlara elle girebilirsin.", "rows": default_rows(exam_type)}


@router.post("", response_model=ExamResult)
async def save_exam(payload: ExamIn, user: dict = Depends(student_only)):
    rows = [analyse(r.model_dump(), payload.exam_type) for r in payload.rows]
    total_net = round(sum(r["net"] for r in rows), 2)
    total_questions = sum(r["max_questions"] or (r["correct"] + r["wrong"] + r["blank"]) for r in rows)
    ranked = sorted(rows, key=lambda r: (r["rate"], -r["max_questions"]))
    weakest, strongest = ranked[0], ranked[-1]
    heavy = max(rows, key=lambda r: (r["max_questions"] - r["correct"]) )
    advice = [
        f"{weakest['subject']}: {weakest['max_questions'] or '—'} sorudan {weakest['net']} net (%{weakest['rate']}). {weakest['advice']}",
        f"En büyük net potansiyeli {heavy['subject']} dersinde: {heavy['max_questions'] - heavy['correct']} soru daha kazanılabilir.",
        f"{strongest['subject']} dersinde %{strongest['rate']} başarı ile güçlüsün; haftalık tarama testleriyle koru.",
    ]
    coef = 3.3 if payload.exam_type == "TYT" else 3.0
    doc = {"id": str(uuid.uuid4()), "user_id": user["id"], "name": payload.name, "exam_type": payload.exam_type, "total_questions": total_questions, "total_net": total_net,
           "estimated_score": round(min(500.0, 100 + total_net * coef), 1), "rows": rows, "weakest": weakest["subject"], "strongest": strongest["subject"], "advice": advice, "created_at": now_iso()}
    await db.exam_results.insert_one(doc)
    await award(user["id"], POINTS["exam"], "exam")
    doc.pop("_id", None)
    doc.pop("user_id")
    return doc


@router.get("", response_model=list[ExamResult])
async def list_exams(user: dict = Depends(student_only)):
    docs = await db.exam_results.find({"user_id": user["id"]}, {"_id": 0, "user_id": 0}).sort("created_at", -1).to_list(50)
    return [{"exam_type": "TYT", "total_questions": 0, **d} for d in docs]
